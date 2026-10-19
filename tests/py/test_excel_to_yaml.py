import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))

import excel_to_yaml as e2y  # noqa: E402
from excel_to_yaml import ExcelError, convert, valider_echeances  # noqa: E402

CASES = json.loads((ROOT / "tests/fixtures/echeances-cases.json").read_text(encoding="utf-8"))

H_MENSUEL = ["Type", "Catégorie", "Libellé", "Montant", "Revalorisé", "Fin"]
H_ANNUEL = ["Type", "Catégorie", "Libellé", "Échéances", "Revalorisé", "Début", "Fin"]
H_EXC = ["Type", "Catégorie", "Libellé", "Échéances"]


def make_xlsx(path, mensuel=None, annuel=None, exc=None, h_annuel=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Situation"
    ws.append(["Poste", "Montant"])
    ws.append(["Compte courant", 350])
    ws.append(["Épargne", 4000])
    for name, header, rows in [
        ("Mensuel", H_MENSUEL, mensuel if mensuel is not None else [
            ["revenu", "Salaire", "Paie", 3000, "oui", None],
            ["dépense", "Logement", "Loyer", 800, "non", "2030-12"],
        ]),
        ("Annuel", h_annuel or H_ANNUEL, annuel if annuel is not None else [
            ["dépense", "Impôts", "Taxe", "6:500, 10:500", "oui", "2027-01", None],
            ["revenu", "Prime", "Prime Leader", "3-300", "non", None, "2031-06"],
        ]),
        ("Exceptionnel", H_EXC, exc if exc is not None else [
            ["revenu", "Divers", "Héritage", "2029-03:25000"],
            ["dépense", "Voiture", "Achat", "2026-3-300"],
        ]),
    ]:
        w = wb.create_sheet(name)
        w.append(header)
        for r in rows:
            w.append(r)
    wb.save(path)
    return path


def test_conversion_nominale(tmp_path):
    d = convert(make_xlsx(tmp_path / "b.xlsx"))
    assert d["situation"] == {"compte_courant": 350, "epargne": 4000}
    assert d["mensuel"] == [
        {"type": "revenu", "categorie": "Salaire", "libelle": "Paie", "montant": 3000, "revalorise": True, "fin": None},
        {"type": "depense", "categorie": "Logement", "libelle": "Loyer", "montant": 800, "revalorise": False, "fin": "2030-12"},
    ]
    assert d["annuel"] == [
        {"type": "depense", "categorie": "Impôts", "libelle": "Taxe", "echeances": "6:500, 10:500", "revalorise": True, "debut": "2027-01", "fin": None},
        {"type": "revenu", "categorie": "Prime", "libelle": "Prime Leader", "echeances": "3-300", "revalorise": False, "debut": None, "fin": "2031-06"},
    ]
    assert d["exceptionnel"] == [
        {"type": "revenu", "categorie": "Divers", "libelle": "Héritage", "echeances": "2029-03:25000"},
        {"type": "depense", "categorie": "Voiture", "libelle": "Achat", "echeances": "2026-3-300"},
    ]


@pytest.mark.parametrize("texte", [c["texte"] for c in CASES["valides"]])
def test_echeances_valides(texte):
    valider_echeances(texte)


@pytest.mark.parametrize("texte", CASES["invalides"])
def test_echeances_invalides(texte):
    with pytest.raises(ValueError):
        valider_echeances(texte)


def test_echeance_3_300_dans_annuel(tmp_path):
    p = make_xlsx(tmp_path / "b.xlsx", annuel=[["dépense", "A", "B", "3-300", "non", None, None]])
    assert convert(p)["annuel"][0]["echeances"] == "3-300"


def test_echeance_numerique_et_dates_excel(tmp_path):
    import datetime
    p = make_xlsx(
        tmp_path / "b.xlsx",
        annuel=[["Depense", "A", "B", 1237.5, "Oui", datetime.datetime(2027, 5, 1), None]],
    )
    a = convert(p)["annuel"][0]
    assert a["echeances"] == "1237.5"
    assert a["debut"] == "2027-05"
    assert a["type"] == "depense"


def test_entete_manquant(tmp_path):
    p = make_xlsx(tmp_path / "b.xlsx", h_annuel=["Type", "Catégorie", "Libellé", "Échéances", "Revalorisé", "Début"])
    with pytest.raises(ExcelError, match="Annuel"):
        convert(p)


def test_echeance_invalide_ligne_3(tmp_path):
    p = make_xlsx(tmp_path / "b.xlsx", annuel=[
        ["dépense", "A", "B", "3-300", "non", None, None],
        ["dépense", "A", "C", "13:500", "non", None, None],
    ])
    with pytest.raises(ExcelError) as ei:
        convert(p)
    assert "Annuel" in str(ei.value) and "ligne 3" in str(ei.value)


def test_fin_avant_debut(tmp_path):
    p = make_xlsx(tmp_path / "b.xlsx", annuel=[["dépense", "A", "B", "3-300", "non", "2030-01", "2029-12"]])
    with pytest.raises(ExcelError, match="Annuel"):
        convert(p)


def test_cli_echec_preserve_sortie(tmp_path):
    p = make_xlsx(tmp_path / "b.xlsx", annuel=[["dépense", "A", "B", "abc", "non", None, None]])
    sortie = tmp_path / "out.yaml"
    sortie.write_text("ancien: contenu\n", encoding="utf-8")
    r = subprocess.run(
        [sys.executable, str(ROOT / "scripts/excel_to_yaml.py"), str(p), "--sortie", str(sortie)],
        capture_output=True, text=True,
    )
    assert r.returncode == 1
    assert "Annuel" in r.stderr
    assert sortie.read_text(encoding="utf-8") == "ancien: contenu\n"


def test_cli_succes(tmp_path):
    p = make_xlsx(tmp_path / "b.xlsx")
    sortie = tmp_path / "sub" / "out.yaml"
    r = subprocess.run(
        [sys.executable, str(ROOT / "scripts/excel_to_yaml.py"), str(p), "--sortie", str(sortie)],
        capture_output=True, text=True,
    )
    assert r.returncode == 0, r.stderr
    import yaml
    assert yaml.safe_load(sortie.read_text(encoding="utf-8")) == convert(p)


def test_vrai_budget():
    real = ROOT / "budget.xlsx"
    if not real.exists():
        pytest.skip("budget.xlsx absent")
    d = convert(real)
    wb = openpyxl.load_workbook(real, read_only=True, data_only=True)
    n = sum(1 for r in list(wb["Mensuel"].iter_rows(min_row=2, values_only=True)) if any(c not in (None, "") for c in r))
    assert len(d["mensuel"]) == n
    leader = [a for a in d["annuel"] if a["libelle"] == "Prime Leader"]
    assert leader and leader[0]["debut"] == "2029-01"
