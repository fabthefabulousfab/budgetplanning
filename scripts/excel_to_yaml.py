"""Convertit budget.xlsx en data/donnees-courantes.yaml."""
import argparse
import datetime
import os
import re
import sys
import tempfile
import unicodedata
from pathlib import Path

import openpyxl
import yaml

H_MENSUEL = ["Type", "Catégorie", "Libellé", "Montant", "Revalorisé", "Fin"]
H_ANNUEL = ["Type", "Catégorie", "Libellé", "Échéances", "Revalorisé", "Début", "Fin"]
H_EXC = ["Type", "Catégorie", "Libellé", "Échéances"]

_MENSUEL_RE = re.compile(r"^\d+(\.\d+)?$")
_ITEM_RE = re.compile(r"^(?:(\d{4})-)?(\d{1,2})[-:](\d+(?:\.\d+)?)$")
_MOIS_RE = re.compile(r"^(\d{4})-(\d{1,2})$")


class ExcelError(Exception):
    pass


def valider_echeances(texte):
    """Même grammaire que js/schedule.js. Lève ValueError."""
    t = (texte or "").strip()
    if not t:
        raise ValueError("échéances vides")
    if _MENSUEL_RE.match(t):
        if float(t) <= 0:
            raise ValueError(f"montant invalide : {t}")
        return
    datees = []
    for item in t.split(","):
        m = _ITEM_RE.match(item.strip())
        if not m:
            raise ValueError(f"échéance invalide : « {item.strip()} »")
        if not 1 <= int(m.group(2)) <= 12:
            raise ValueError(f"mois invalide : « {item.strip()} »")
        if float(m.group(3)) <= 0:
            raise ValueError(f"montant invalide : « {item.strip()} »")
        datees.append(m.group(1) is not None)
    if any(datees) and not all(datees):
        raise ValueError("mélange d'échéances datées et non datées")


def _norm(s):
    s = unicodedata.normalize("NFD", str(s).strip().lower())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def _vide(v):
    return v is None or (isinstance(v, str) and not v.strip())


def _rows(wb, nom, entetes):
    if nom not in wb.sheetnames:
        raise ExcelError(f"Feuille {nom} : feuille absente")
    it = wb[nom].iter_rows(values_only=True)
    head = next(it, None)
    head = [_norm(c) if c is not None else "" for c in (head or [])]
    manquants = [h for h in entetes if _norm(h) not in head]
    if manquants:
        raise ExcelError(f"Feuille {nom}, ligne 1 : en-tête(s) manquant(s) : {', '.join(manquants)}")
    idx = {h: head.index(_norm(h)) for h in entetes}
    for n, row in enumerate(it, start=2):
        if all(_vide(c) for c in row):
            continue
        yield n, {h: (row[i] if i < len(row) else None) for h, i in idx.items()}


def _type(nom, n, v):
    k = _norm(v) if v is not None else ""
    if k == "revenu":
        return "revenu"
    if k == "depense":
        return "depense"
    raise ExcelError(f"Feuille {nom}, ligne {n} : Type invalide « {v} » (revenu ou dépense)")


def _texte(nom, n, champ, v):
    if _vide(v):
        raise ExcelError(f"Feuille {nom}, ligne {n} : {champ} manquant")
    return str(v).strip()


def _bool(nom, n, v):
    k = _norm(v) if v is not None else ""
    if k == "oui":
        return True
    if k == "non" or k == "":
        return False
    raise ExcelError(f"Feuille {nom}, ligne {n} : Revalorisé invalide « {v} » (oui ou non)")


def _mois(nom, n, champ, v):
    if _vide(v):
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return f"{v.year:04d}-{v.month:02d}"
    s = str(v).strip()
    m = _MOIS_RE.match(s)
    if not m or not 1 <= int(m.group(2)) <= 12:
        raise ExcelError(f"Feuille {nom}, ligne {n} : {champ} invalide « {s} » (AAAA-MM attendu)")
    return f"{m.group(1)}-{int(m.group(2)):02d}"


def _nombre(nom, n, champ, v):
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        raise ExcelError(f"Feuille {nom}, ligne {n} : {champ} doit être un nombre")
    return v


def _montant_positif(nom, n, champ, v):
    v = _nombre(nom, n, champ, v)
    if v <= 0:
        raise ExcelError(f"Feuille {nom}, ligne {n} : {champ} doit être supérieur à 0")
    return v


def _echeances(nom, n, v):
    if _vide(v):
        raise ExcelError(f"Feuille {nom}, ligne {n} : Échéances manquantes")
    s = repr(v) if isinstance(v, float) else str(v)
    s = s.strip()
    try:
        valider_echeances(s)
    except ValueError as err:
        raise ExcelError(f"Feuille {nom}, ligne {n} : {err}") from None
    return s


def _periode(nom, n, debut, fin):
    if debut and fin and fin < debut:
        raise ExcelError(f"Feuille {nom}, ligne {n} : Fin ({fin}) antérieure à Début ({debut})")


def convert(xlsx_path):
    wb = openpyxl.load_workbook(Path(xlsx_path), read_only=True, data_only=True)
    try:
        situation = {}
        for n, r in _rows(wb, "Situation", ["Poste", "Montant"]):
            k = _norm(r["Poste"])
            cle = {"compte courant": "compte_courant", "epargne": "epargne"}.get(k)
            if cle:
                situation[cle] = _nombre("Situation", n, "Montant", r["Montant"])
        for cle, label in (("compte_courant", "Compte courant"), ("epargne", "Épargne")):
            if cle not in situation:
                raise ExcelError(f"Feuille Situation : ligne « {label} » absente")

        mensuel = []
        for n, r in _rows(wb, "Mensuel", H_MENSUEL):
            mensuel.append({
                "type": _type("Mensuel", n, r["Type"]),
                "categorie": _texte("Mensuel", n, "Catégorie", r["Catégorie"]),
                "libelle": _texte("Mensuel", n, "Libellé", r["Libellé"]),
                "montant": _montant_positif("Mensuel", n, "Montant", r["Montant"]),
                "revalorise": _bool("Mensuel", n, r["Revalorisé"]),
                "fin": _mois("Mensuel", n, "Fin", r["Fin"]),
            })

        annuel = []
        for n, r in _rows(wb, "Annuel", H_ANNUEL):
            debut = _mois("Annuel", n, "Début", r["Début"])
            fin = _mois("Annuel", n, "Fin", r["Fin"])
            _periode("Annuel", n, debut, fin)
            annuel.append({
                "type": _type("Annuel", n, r["Type"]),
                "categorie": _texte("Annuel", n, "Catégorie", r["Catégorie"]),
                "libelle": _texte("Annuel", n, "Libellé", r["Libellé"]),
                "echeances": _echeances("Annuel", n, r["Échéances"]),
                "revalorise": _bool("Annuel", n, r["Revalorisé"]),
                "debut": debut,
                "fin": fin,
            })

        exceptionnel = []
        for n, r in _rows(wb, "Exceptionnel", H_EXC):
            exceptionnel.append({
                "type": _type("Exceptionnel", n, r["Type"]),
                "categorie": _texte("Exceptionnel", n, "Catégorie", r["Catégorie"]),
                "libelle": _texte("Exceptionnel", n, "Libellé", r["Libellé"]),
                "echeances": _echeances("Exceptionnel", n, r["Échéances"]),
            })
    finally:
        wb.close()
    return {"situation": situation, "mensuel": mensuel, "annuel": annuel, "exceptionnel": exceptionnel}


def ecrire_atomique(data, sortie):
    sortie = Path(sortie)
    sortie.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=sortie.parent, prefix=sortie.name + ".", suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            yaml.safe_dump(data, f, allow_unicode=True, sort_keys=False)
        os.replace(tmp, sortie)
    except BaseException:
        if os.path.exists(tmp):
            os.unlink(tmp)
        raise


def main(argv=None):
    ap = argparse.ArgumentParser(description="Convertit budget.xlsx en YAML.")
    ap.add_argument("xlsx")
    ap.add_argument("--sortie", default="data/donnees-courantes.yaml")
    args = ap.parse_args(argv)
    try:
        data = convert(Path(args.xlsx))
        ecrire_atomique(data, args.sortie)
    except (ExcelError, OSError) as err:
        print(f"Erreur : {err}", file=sys.stderr)
        return 1
    print(
        f"{args.sortie} écrit : {len(data['mensuel'])} mensuel, "
        f"{len(data['annuel'])} annuel, {len(data['exceptionnel'])} exceptionnel"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
